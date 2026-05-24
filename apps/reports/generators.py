import io
import os
from datetime import date
from decimal import Decimal

from django.conf import settings

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

# ===== Регистрация шрифта =====
FONTS_DIR = settings.BASE_DIR / "static" / "fonts"

pdfmetrics.registerFont(TTFont("AppFont",      str(FONTS_DIR / "Roboto-Regular.ttf")))
pdfmetrics.registerFont(TTFont("AppFont-Bold", str(FONTS_DIR / "Roboto-Bold.ttf")))
pdfmetrics.registerFontFamily("AppFont", normal="AppFont", bold="AppFont-Bold")


# ===== DOCX Generators =====

def _add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_table_header(table, headers: list[str], bg_color="1F5C99"):
    row = table.rows[0]
    for i, header in enumerate(headers):
        cell = row.cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell._tc.get_or_add_tcPr()


def generate_stock_report_docx(data: list[dict]) -> io.BytesIO:
    doc = Document()
    _add_heading(doc, "Отчёт по наличию лекарств на складе", level=1)
    _add_heading(doc, f"Дата: {date.today().strftime('%d.%m.%Y')}", level=2)
    doc.add_paragraph()

    for group_data in data:
        doc.add_heading(group_data["group"], level=2)

        if not group_data["drugs"]:
            doc.add_paragraph("Нет препаратов в данной группе")
            continue

        table = doc.add_table(rows=1 + len(group_data["drugs"]) + 1, cols=3)
        table.style = "Table Grid"

        headers = ["Наименование препарата", "Остаток", "Ед. изм."]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        for j, drug in enumerate(group_data["drugs"], 1):
            row = table.rows[j]
            row.cells[0].text = drug["name"]
            row.cells[1].text = str(drug["quantity"])
            row.cells[2].text = "шт."
            row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        total_row = table.rows[-1]
        total_row.cells[0].text = "Итого по группе:"
        total_row.cells[0].paragraphs[0].runs[0].bold = True
        total_row.cells[1].text = str(group_data["total"])
        total_row.cells[1].paragraphs[0].runs[0].bold = True
        total_row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_sales_report_docx(data: list[dict], date_from, date_to) -> io.BytesIO:
    doc = Document()
    _add_heading(doc, "Отчёт по продажам лекарств по группам", level=1)
    _add_heading(doc, f"Период: {date_from.strftime('%d.%m.%Y')} — {date_to.strftime('%d.%m.%Y')}", level=2)
    doc.add_paragraph()

    grand_total_amount = Decimal("0")

    for group_data in data:
        doc.add_heading(group_data["group"], level=2)

        table = doc.add_table(rows=1 + len(group_data["drugs"]) + 1, cols=4)
        table.style = "Table Grid"

        headers = ["Наименование препарата", "Кол-во", "Сумма (руб.)", ""]
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        for j, drug in enumerate(group_data["drugs"], 1):
            row = table.rows[j]
            row.cells[0].text = drug["drug__name"]
            row.cells[1].text = str(drug["quantity"])
            row.cells[2].text = f"{drug['amount']:.2f}"
            row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        total_row = table.rows[-1]
        total_row.cells[0].text = "Итого по группе:"
        total_row.cells[0].paragraphs[0].runs[0].bold = True
        total_row.cells[1].text = str(group_data["total_quantity"])
        total_row.cells[1].paragraphs[0].runs[0].bold = True
        total_row.cells[2].text = f"{group_data['total_amount']:.2f}"
        total_row.cells[2].paragraphs[0].runs[0].bold = True
        total_row.cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

        grand_total_amount += group_data["total_amount"] or Decimal("0")
        doc.add_paragraph()

    p = doc.add_paragraph()
    run = p.add_run(f"ИТОГО ПО ВСЕМ ГРУППАМ: {grand_total_amount:.2f} руб.")
    run.bold = True
    run.font.size = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_invoice_docx(data: dict) -> io.BytesIO:
    doc = Document()
    _add_heading(doc, "СЧЁТ-ФАКТУРА", level=1)
    doc.add_paragraph()

    info_table = doc.add_table(rows=4, cols=2)
    info_table.style = "Table Grid"
    rows_data = [
        ("Номер документа:", data["document_number"]),
        ("Дата:",            data["document_date"].strftime("%d.%m.%Y")),
        ("Поставщик:",       data["supplier"]),
        ("ИНН поставщика:",  data["supplier_inn"] or ""),
    ]
    for i, (label, value) in enumerate(rows_data):
        info_table.rows[i].cells[0].text = label
        info_table.rows[i].cells[0].paragraphs[0].runs[0].bold = True
        info_table.rows[i].cells[1].text = value

    doc.add_paragraph()

    table = doc.add_table(rows=1 + len(data["items"]) + 1, cols=6)
    table.style = "Table Grid"

    headers = ["Наименование", "Серия", "Срок годности", "Кол-во", "Цена", "Сумма"]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for j, item in enumerate(data["items"], 1):
        row = table.rows[j]
        row.cells[0].text = item["name"]
        row.cells[1].text = item["series"]
        row.cells[2].text = item["expiry_date"].strftime("%d.%m.%Y")
        row.cells[3].text = str(item["quantity"])
        row.cells[4].text = f"{item['price']:.2f}"
        row.cells[5].text = f"{item['total']:.2f}"
        for k in range(3, 6):
            row.cells[k].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    total_row = table.rows[-1]
    total_row.cells[0].text = "ИТОГО:"
    total_row.cells[0].paragraphs[0].runs[0].bold = True
    total_row.cells[5].text = f"{data['grand_total']:.2f}"
    total_row.cells[5].paragraphs[0].runs[0].bold = True
    total_row.cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph()
    doc.add_paragraph(f"Принял: {data['created_by']}")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ===== PDF Generators =====

def _get_pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="RuTitle",
        fontName="AppFont-Bold",   
        fontSize=16,
        spaceAfter=12,
        alignment=1,
    ))
    styles.add(ParagraphStyle(
        name="RuHeading",
        fontName="AppFont-Bold",   
        fontSize=12,
        spaceAfter=8,
    ))
    styles.add(ParagraphStyle(
        name="RuNormal",
        fontName="AppFont",        
        fontSize=10,
    ))
    styles.add(ParagraphStyle(
        name="RuBold",
        fontName="AppFont-Bold",   
        fontSize=10,
    ))
    return styles


def _get_table_style():
    return TableStyle([
        ("BACKGROUND",    (0, 0),  (-1, 0),  colors.HexColor("#1F5C99")),
        ("TEXTCOLOR",     (0, 0),  (-1, 0),  colors.white),
        ("FONTNAME",      (0, 0),  (-1, 0),  "AppFont-Bold"),
        ("FONTNAME",      (0, 1),  (-1, -1), "AppFont"),
        ("FONTSIZE",      (0, 0),  (-1, -1), 10),
        ("ALIGN",         (0, 0),  (-1, 0),  "CENTER"),
        ("GRID",          (0, 0),  (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS",(0, 1),  (-1, -2), [colors.white, colors.HexColor("#F0F4FA")]),
        ("BACKGROUND",    (0, -1), (-1, -1), colors.HexColor("#D5E8F0")),
        ("FONTNAME",      (0, -1), (-1, -1), "AppFont-Bold"),
        ("ALIGN",         (1, 1),  (-1, -1), "CENTER"),
        ("VALIGN",        (0, 0),  (-1, -1), "MIDDLE"),
        ("TOPPADDING",    (0, 0),  (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0),  (-1, -1), 6),
    ])


def generate_stock_report_pdf(data: list[dict]) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = _get_pdf_styles()
    story = []

    story.append(Paragraph("Отчёт по наличию лекарств на складе", styles["RuTitle"]))
    story.append(Paragraph(f"Дата: {date.today().strftime('%d.%m.%Y')}", styles["RuNormal"]))
    story.append(Spacer(1, 0.5*cm))

    for group_data in data:
        story.append(Paragraph(group_data["group"], styles["RuHeading"]))

        if not group_data["drugs"]:
            story.append(Paragraph("Нет препаратов в данной группе", styles["RuNormal"]))
            continue

        table_data = [["Наименование препарата", "Остаток", "Ед. изм."]]
        for drug in group_data["drugs"]:
            table_data.append([drug["name"], str(drug["quantity"]), "шт."])
        table_data.append(["Итого по группе:", str(group_data["total"]), ""])

        t = Table(table_data, colWidths=[10*cm, 3*cm, 3*cm])
        t.setStyle(_get_table_style())
        story.append(t)
        story.append(Spacer(1, 0.5*cm))

    doc.build(story)
    buffer.seek(0)
    return buffer


def generate_sales_report_pdf(data: list[dict], date_from, date_to) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = _get_pdf_styles()
    story = []

    story.append(Paragraph("Отчёт по продажам лекарств по группам", styles["RuTitle"]))
    story.append(Paragraph(
        f"Период: {date_from.strftime('%d.%m.%Y')} — {date_to.strftime('%d.%m.%Y')}",
        styles["RuNormal"],
    ))
    story.append(Spacer(1, 0.5*cm))

    grand_total = Decimal("0")

    for group_data in data:
        story.append(Paragraph(group_data["group"], styles["RuHeading"]))

        table_data = [["Наименование препарата", "Кол-во", "Сумма (руб.)"]]
        for drug in group_data["drugs"]:
            table_data.append([
                drug["drug__name"],
                str(drug["quantity"]),
                f"{drug['amount']:.2f}",
            ])
        table_data.append([
            "Итого по группе:",
            str(group_data["total_quantity"]),
            f"{group_data['total_amount']:.2f}",
        ])

        t = Table(table_data, colWidths=[9*cm, 3*cm, 4*cm])
        t.setStyle(_get_table_style())
        story.append(t)
        story.append(Spacer(1, 0.5*cm))
        grand_total += group_data["total_amount"] or Decimal("0")

    story.append(Paragraph(
        f"ИТОГО ПО ВСЕМ ГРУППАМ: {grand_total:.2f} руб.",
        styles["RuBold"],
    ))

    doc.build(story)
    buffer.seek(0)
    return buffer


def generate_invoice_pdf(data: dict) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, rightMargin=2*cm, leftMargin=2*cm,
                            topMargin=2*cm, bottomMargin=2*cm)
    styles = _get_pdf_styles()
    story = []

    story.append(Paragraph("СЧЁТ-ФАКТУРА", styles["RuTitle"]))
    story.append(Spacer(1, 0.3*cm))

    info_data = [
        ["Номер документа:", data["document_number"]],
        ["Дата:", data["document_date"].strftime("%d.%m.%Y")],
        ["Поставщик:", data["supplier"]],
        ["ИНН поставщика:", data["supplier_inn"] or ""],
    ]
    info_table = Table(info_data, colWidths=[5*cm, 11*cm])
    info_table.setStyle(TableStyle([
        ("FONTNAME",      (0, 0), (0, -1),  "AppFont-Bold"),   # ← кириллица
        ("FONTNAME",      (1, 0), (1, -1),  "AppFont"),        # ← кириллица
        ("FONTSIZE",      (0, 0), (-1, -1), 10),
        ("GRID",          (0, 0), (-1, -1), 0.5, colors.grey),
        ("TOPPADDING",    (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(info_table)
    story.append(Spacer(1, 0.5*cm))

    table_data = [["Наименование", "Серия", "Срок годн.", "Кол-во", "Цена", "Сумма"]]
    for item in data["items"]:
        table_data.append([
            item["name"],
            item["series"],
            item["expiry_date"].strftime("%d.%m.%Y"),
            str(item["quantity"]),
            f"{item['price']:.2f}",
            f"{item['total']:.2f}",
        ])
    table_data.append(["ИТОГО:", "", "", "", "", f"{data['grand_total']:.2f}"])

    t = Table(table_data, colWidths=[5*cm, 2*cm, 2.5*cm, 1.5*cm, 2*cm, 3*cm])
    t.setStyle(_get_table_style())
    story.append(t)
    story.append(Spacer(1, 0.5*cm))
    story.append(Paragraph(f"Принял: {data['created_by']}", styles["RuNormal"]))

    doc.build(story)
    buffer.seek(0)
    return buffer